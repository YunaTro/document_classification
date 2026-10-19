"""Адаптированные из старого проекта загрузчики PDF, DOCX и изображений.

Контракт: список секций с текстом, номером страницы и OCR-метаданными.
Никакой индексации или классификации здесь нет.
"""
from pathlib import Path

import fitz
from docx import Document
from PIL import Image

from app.config import settings
from app.ocr_utils import ocr_image_with_confidence

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\x00", " ").replace("\t", " ").replace("\r", "\n")
    return "\n".join(line.strip() for line in text.split("\n") if line.strip())


def load_pdf(path: str) -> list[dict]:
    rows: list[dict] = []
    with fitz.open(path) as pdf:
        for number, page in enumerate(pdf, start=1):
            text = clean_text(page.get_text("text"))
            has_ocr = False
            ocr_confidence = None
            if len(text) < settings.ocr_min_text_chars and settings.enable_ocr:
                pix = page.get_pixmap(dpi=160, alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                recognized, ocr_confidence = ocr_image_with_confidence(image)
                if recognized.strip():
                    text = clean_text(recognized)
                has_ocr = True
            rows.append({
                "section_type": "page", "page_number": number,
                "text": text, "has_ocr": has_ocr, "ocr_confidence": ocr_confidence,
            })
    return rows


def load_docx(path: str) -> list[dict]:
    doc = Document(path)
    rows: list[dict] = []
    # Сохраняем исходную идею старого проекта: абзацы и текст таблиц.
    paragraphs = [clean_text(p.text) for p in doc.paragraphs if clean_text(p.text)]
    if paragraphs:
        rows.append({
            "section_type": "paragraphs", "page_number": None,
            "text": "\n".join(paragraphs), "has_ocr": False, "ocr_confidence": None,
        })
    for table in doc.tables:
        lines = [" | ".join(clean_text(cell.text) for cell in row.cells) for row in table.rows]
        text = clean_text("\n".join(lines))
        if text:
            rows.append({
                "section_type": "table", "page_number": None,
                "text": text, "has_ocr": False, "ocr_confidence": None,
            })
    return rows


def load_image(path: str) -> list[dict]:
    with Image.open(path) as img:
        text, confidence = ocr_image_with_confidence(img.convert("RGB"))
    return [{
        "section_type": "image", "page_number": 1,
        "text": clean_text(text), "has_ocr": settings.enable_ocr,
        "ocr_confidence": confidence,
    }]


def load_any(path: str) -> list[dict]:
    extension = Path(path).suffix.lower()
    if extension == ".pdf":
        return load_pdf(path)
    if extension == ".docx":
        return load_docx(path)
    if extension in SUPPORTED_EXTENSIONS:
        return load_image(path)
    raise ValueError(f"Неподдерживаемый формат: {extension}")
