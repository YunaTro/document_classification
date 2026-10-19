
"""Генерация синтетических документов для обучения."""

import argparse
from pathlib import Path
import random

from docx import Document

from app.document_classes import CLASSES


TOPICS = [
    "поставка офисной бумаги",
    "обслуживание серверного оборудования",
    "ремонт помещений",
    "организация корпоративного обучения",
    "обновление программного обеспечения",
    "закупка мебели",
    "техническая поддержка сотрудников",
    "обслуживание вентиляции",
    "проверка пожарной безопасности",
    "перевозка оборудования",
    "поставка расходных материалов",
    "настройка локальной сети",
    "инвентаризация склада",
    "ремонт оргтехники",
    "доставка канцелярских товаров",
    "разработка рекламных материалов",
    "обслуживание транспортных средств",
    "установка систем освещения",
    "проведение инструктажа",
    "плановое обслуживание здания",
]

COMPANIES = [
    "ООО «Север»",
    "ООО «Альтаир»",
    "АО «Маяк»",
    "ООО «Вектор»",
    "ООО «Меридиан»",
    "АО «Горизонт»",
]

HEADINGS = {
    "contract": [
        "ДОГОВОР",
        "СОГЛАШЕНИЕ О СОТРУДНИЧЕСТВЕ",
        "ДОГОВОР ОКАЗАНИЯ УСЛУГ",
    ],
    "invoice": [
        "СЧЁТ НА ОПЛАТУ",
        "СЧЕТ",
        "СЧЁТ ЗА ОКАЗАННЫЕ УСЛУГИ",
    ],
    "act": [
        "АКТ ПРИЁМКИ",
        "АКТ ВЫПОЛНЕННЫХ РАБОТ",
        "АКТ СДАЧИ-ПРИЁМКИ",
    ],
    "application": [
        "ЗАЯВЛЕНИЕ",
        "ОБРАЩЕНИЕ СОТРУДНИКА",
        "ЗАЯВЛЕНИЕ НА СОГЛАСОВАНИЕ",
    ],
    "memo": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "ВНУТРЕННЯЯ ЗАПИСКА",
        "ДОКЛАДНАЯ ЗАПИСКА",
    ],
    "report": [
        "ОТЧЁТ",
        "ОТЧЕТ О ПРОДЕЛАННОЙ РАБОТЕ",
        "ОТЧЁТ О РЕЗУЛЬТАТАХ",
    ],
}


def make_paragraphs(
    label: str,
    topic: str,
    company: str,
    number: int,
    amount: int,
    family: int,
    copy: int,
) -> list[str]:

    subject = topic.capitalize()
    date = (
        f"{(family % 25) + 1:02d}."
        f"{(family % 9) + 1:02d}.2026"
    )

    intro = (
        f"Документ № {number} от {date}. "
        f"Организация: {company}."
    )

    if label == "contract":
        body = [
            f"Заказчик и исполнитель согласовали "
            f"предмет взаимодействия: {topic}.",
            f"Исполнитель обязуется обеспечить {topic}, "
            "а заказчик — принять результат и оплатить его.",
            f"Согласованная стоимость составляет "
            f"{amount:,} рублей. Срок выполнения "
            "определяется календарным планом.",
            "Стороны несут ответственность за нарушение "
            "условий и подтверждают соглашение подписями.",
        ]

    elif label == "invoice":
        body = [
            f"Плательщик: "
            f"{COMPANIES[(family + 2) % len(COMPANIES)]}. "
            f"Получатель: {company}.",
            f"Наименование позиции: {subject}. "
            f"Количество: {copy + 1} усл. ед.",
            f"Итого к оплате: {amount:,} рублей. "
            f"Назначение платежа: оплата за {topic}.",
            "Просим перечислить сумму по указанным "
            "реквизитам. После оплаты направить подтверждение.",
        ]

    elif label == "act":
        body = [
            f"Комиссия проверила результат работ "
            f"по направлению: {topic}.",
            f"Выполнение подтверждено ответственными "
            f"лицами {company}. Объём работ проверен "
            "по перечню.",
            f"Общая стоимость принятых работ составляет "
            f"{amount:,} рублей.",
            "Замечания по комплектности и качеству: "
            "отсутствуют. Результат передан заказчику.",
        ]

    elif label == "application":
        body = [
            f"Руководителю {company} от сотрудника "
            "административного подразделения.",
            "Прошу рассмотреть мою просьбу: "
            "согласовать участие в мероприятиях, "
            f"связанных с направлением «{topic}».",
            "Прошу предоставить разрешение на выполнение "
            "запланированных действий в рабочее время.",
            f"Необходимость обращения обусловлена планом "
            f"подразделения на {date}. "
            "Подпись заявителя прилагается.",
        ]

    elif label == "memo":
        body = [
            f"Кому: руководителю {company}. "
            "От кого: координатора проекта.",
            "Сообщаю о необходимости дополнительных "
            f"организационных мер по вопросу: {topic}.",
            f"Предлагаю согласовать финансирование "
            f"в размере {amount:,} рублей "
            "и назначить ответственного.",
            "Прошу рассмотреть предложение на рабочем "
            "совещании и сообщить принятое решение.",
        ]

    elif label == "report":
        body = [
            f"Предмет отчётного периода: {topic}. "
            f"Подразделение: {company}.",
            "За отчётный период выполнен анализ исходной "
            "ситуации, составлен план и организованы мероприятия.",
            f"Проведено {copy + family % 5 + 2} "
            "контрольных действий; оценены сроки, "
            "затраты и результаты.",
            "По итогам отмечено выполнение основных "
            "показателей. Рекомендовано продолжить наблюдение.",
        ]

    else:
        raise ValueError(label)

    rotation = family % len(body)
    body = body[rotation:] + body[:rotation]

    return [
        intro,
        *body,
        "Ответственный: __________________. "
        "Дата: __________________.",
    ]


def generate(root: Path, *, force: bool = False) -> int:

    targets = [
        root / label /
        f"{label}-tpl{family:02d}__{copy:02d}.docx"
        for label in CLASSES
        for family in range(20)
        for copy in range(2)
    ]

    existing = [p for p in targets if p.exists()]

    if existing and not force:
        raise FileExistsError(
            f"Найдены готовые файлы ({len(existing)}). "
            "Для перезаписи созданных файлов укажи --force."
        )

    for label in CLASSES:
        (root / label).mkdir(
            parents=True, exist_ok=True
        )

        for family in range(20):
            for copy in range(2):
                rng = random.Random(
                    f"2026:{label}:{family}:{copy}"
                )

                company = COMPANIES[
                    rng.randrange(len(COMPANIES))
                ]

                topic = TOPICS[family]
                amount = rng.randint(15, 260) * 1000
                number = (family + 1) * 100 + copy

                document = Document()

                document.add_heading(
                    HEADINGS[label][family % 3],
                    level=0,
                )

                paragraphs = make_paragraphs(
                    label,
                    topic,
                    company,
                    number,
                    amount,
                    family,
                    copy,
                )

                for paragraph in paragraphs:
                    document.add_paragraph(paragraph)

                path = (
                    root / label /
                    f"{label}-tpl{family:02d}__{copy:02d}.docx"
                )

                document.save(path)

    return len(targets)


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--out",
        type=Path,
        default=Path("data/raw"),
    )

    parser.add_argument(
        "--force",
        action="store_true",
    )

    args = parser.parse_args()

    count = generate(args.out, force=args.force)

    print(
        f"Создано {count} учебных DOCX в {args.out}"
    )


if __name__ == "__main__":
    main()
