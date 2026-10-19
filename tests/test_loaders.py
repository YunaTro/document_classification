"""Автономные тесты извлечения текста."""
from dataclasses import replace

import fitz
from docx import Document
from PIL import Image

from app import loaders


def test_clean_text():
    assert loaders.clean_text("  строка  \n\n  другая\tстрока\x00\n") == "строка\nдругая строка"


def test_docx_paragraphs_and_table(tmp_path):
    file = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("ДОГОВОР № 1")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Продавец"
    table.cell(0, 1).text = "Покупатель"
    doc.save(file)
    sections = loaders.load_any(str(file))
    assert len(sections) == 2
    assert sections[0]["section_type"] == "paragraphs"
    assert "ДОГОВОР" in sections[0]["text"]
    assert sections[1]["section_type"] == "table"
    assert "Продавец | Покупатель" in sections[1]["text"]


def test_pdf_text_pages(tmp_path, monkeypatch):
    monkeypatch.setattr(loaders, "settings", replace(loaders.settings, enable_ocr=False))
    path = tmp_path / "report.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 72), "Report: first page contains a summary")
        page = doc.new_page()
        page.insert_text((72, 72), "Report: second page contains results")
        doc.save(str(path))
    sections = loaders.load_any(str(path))
    assert len(sections) == 2
    assert [x["page_number"] for x in sections] == [1, 2]
    assert "summary" in sections[0]["text"]
    assert not sections[1]["has_ocr"]


def test_image_ocr_metadata(tmp_path, monkeypatch):
    path = tmp_path / "scan.png"
    Image.new("RGB", (100, 100), "white").save(path)
    monkeypatch.setattr(loaders, "ocr_image_with_confidence", lambda img: ("СЧЁТ № 9", 93.2))
    sections = loaders.load_any(str(path))
    assert sections[0]["has_ocr"]
    assert sections[0]["ocr_confidence"] == 93.2
    assert sections[0]["text"] == "СЧЁТ № 9"


def test_unsupported_extension(tmp_path):
    path = tmp_path / "file.txt"
    path.write_text("hello")
    try:
        loaders.load_any(str(path))
    except ValueError as error:
        assert "Неподдерживаемый" in str(error)
    else:
        raise AssertionError("Expected ValueError")
